"""
ladder_edits.py -- the ladder keeper's export: every standard code to add to,
or remove from, a node in a Word ladder, as an Excel workbook.

See docs/brief_ladder_edits_export.md. Two sources, both already-signed-off
writer judgment in mh2_seq.db, checked against the ladders as mh2.db now
reads them:

  Add     tag_proposal in 'open' / 'needs_attention', unless the named node
          already carries the standard (landed, not yet reconciled).
  Remove  tag_review with outcome 'incorrect_tag' whose node still carries
          the standard -- see removal_state().

Everything here is derived at call time and nothing is written back: a
download is a snapshot, the ladders are the only thing that closes an edit
(§6, R-H1), and tag_review rows are never modified (§5.2).

Reads mh2.db (read-only connection supplied by the caller) and mh2_seq.db
through review_store. The mh2.db queries below are this module's own, the
same way node_lookup.py owns its one.
"""
from __future__ import annotations

import io
import sqlite3
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from mh2 import node_lookup, review_store
from mh2.coverage import Tag, _collapse_ws, build_tags_by_standard

REMOVED = "removed"
STILL_TAGGED = "still_tagged"
NEEDS_ATTENTION = "needs_attention"

ATTENTION_NOTE = ("Node reworded or removed since this was suggested. The node "
                  "text shown is how it read then; find the node it became, or "
                  "ask the lead.")


@dataclass
class LadderEdit:
    action: str                 # 'Add' | 'Remove'
    ladder_file: str | None
    stem_name: str | None
    seq: int | None             # None when the node is no longer in `nodes`
    concept_skill: str | None
    node_text: str
    standard_id: str
    as_written: str             # the code as it should appear / appears in the cell
    standard_text: str
    why: str | None
    suggested_by: str
    suggested_on: str           # YYYY-MM-DD
    needs_attention: bool


def removal_state(review: dict, node_by_key: dict[str, node_lookup.NodeInfo],
                  tags_by_standard: dict[str, list[Tag]]) -> tuple[str, list[Tag]]:
    """Classify one 'incorrect_tag' review against the current ladders.

    Returns (state, tags), where tags are the standard's tags still on the
    reviewed node -- non-empty only for STILL_TAGGED. Resolved through
    build_tags_by_standard, so a ladder spelling that differs from the
    reviewed standard_id (e.g. '4.NBT.4' for 4.NBT.B.4) still counts as the
    tag being present.
    """
    info = node_by_key.get(review["source_key"])
    if info is None:
        return NEEDS_ATTENTION, []
    on_node = [t for t in tags_by_standard.get(review["standard_id"], [])
               if t.node_id == info.node_id]
    return (STILL_TAGGED, on_node) if on_node else (REMOVED, [])


def _concept_skills(con: sqlite3.Connection) -> dict[str, str]:
    return {node_id: _collapse_ws(cs) for node_id, cs in con.execute(
        "SELECT node_id, concept_skill FROM nodes WHERE concept_skill IS NOT NULL")}


def _standard_texts(con: sqlite3.Connection) -> dict[str, str]:
    return {sid: _collapse_ws(text) for sid, text in con.execute(
        "SELECT standard_id, text FROM standards WHERE text IS NOT NULL")}


def ladders_last_read(con: sqlite3.Connection) -> str | None:
    """When the ladders in this mh2.db were ingested, YYYY-MM-DD. ingest_log
    lives inside mh2.db, so it moves with the database when a rebuild swaps
    it in, unlike rebuild_log.tsv, which can also record failed runs."""
    row = con.execute("SELECT MAX(ts) FROM ingest_log").fetchone()
    return row[0][:10] if row and row[0] else None


def build_ladder_edits(mh2_con: sqlite3.Connection,
                       seq_con: sqlite3.Connection) -> list[LadderEdit]:
    tags_by_standard = build_tags_by_standard(mh2_con)
    node_by_key = node_lookup.build_source_key_lookup(mh2_con)
    stem_names = node_lookup.build_stem_names(mh2_con)
    concepts = _concept_skills(mh2_con)
    std_text = _standard_texts(mh2_con)

    def located(source_key, ladder_file_seen, node_text_seen):
        info = node_by_key.get(source_key)
        if info is None:
            stem_id = source_key.partition(":")[0]
            return dict(ladder_file=ladder_file_seen, stem_name=stem_names.get(stem_id),
                        seq=None, concept_skill=None,
                        node_text=_collapse_ws(node_text_seen or ""),
                        needs_attention=True)
        return dict(ladder_file=info.source_file, stem_name=stem_names.get(info.stem_id),
                    seq=info.seq, concept_skill=concepts.get(info.node_id),
                    node_text=_collapse_ws(info.node_text), needs_attention=False)

    edits: list[LadderEdit] = []

    proposals = sorted(review_store.list_proposals(seq_con, "open")
                       + review_store.list_proposals(seq_con, "needs_attention"),
                       key=lambda p: p["proposal_id"])
    # Duplicate open proposals for one (standard, node) predate the route's
    # open_proposal_exists guard (live mh2_seq.db holds one such pair,
    # #2/#3). They are one instruction to the keeper; the oldest speaks for it.
    seen: set[tuple[str, str]] = set()
    for p in proposals:
        pair = (p["standard_id"], p["source_key"])
        if pair in seen:
            continue
        seen.add(pair)
        info = node_by_key.get(p["source_key"])
        if info is not None and any(t.node_id == info.node_id
                                    for t in tags_by_standard.get(p["standard_id"], [])):
            continue  # landed; the next reconcile closes it
        edits.append(LadderEdit(
            action="Add", standard_id=p["standard_id"], as_written=p["standard_id"],
            standard_text=std_text.get(p["standard_id"], ""), why=p["rationale"],
            suggested_by=p["proposed_by"], suggested_on=p["proposed_at"][:10],
            **located(p["source_key"], p["ladder_file"], p["node_text_seen"])))

    for r in review_store.list_all_tag_reviews(seq_con):
        if r["outcome"] != "incorrect_tag":
            continue
        state, on_node = removal_state(r, node_by_key, tags_by_standard)
        if state == REMOVED:
            continue
        written = sorted({t.raw_code for t in on_node}) or [r["standard_id"]]
        edits.append(LadderEdit(
            action="Remove", standard_id=r["standard_id"], as_written=", ".join(written),
            standard_text=std_text.get(r["standard_id"], ""), why=r["note"],
            suggested_by=r["reviewed_by"], suggested_on=r["reviewed_at"][:10],
            **located(r["source_key"], r["ladder_file_seen"], r["node_text_seen"])))

    edits.sort(key=lambda e: (e.ladder_file or "", e.seq is None, e.seq or 0,
                              e.action != "Remove", e.standard_id))
    return edits


# ------------------------------------------------------------------- workbook

_FONT = "Arial"
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_ADD_FILL = PatternFill("solid", fgColor="E2EFDA")
_REMOVE_FILL = PatternFill("solid", fgColor="FCE4D6")
_ATTENTION_FILL = PatternFill("solid", fgColor="FFF2CC")
_THIN = Side(style="thin", color="BFBFBF")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_COLUMNS = [  # (header, width)
    ("Done", 7), ("Ladder file", 34), ("Stem", 16), ("Node #", 8),
    ("Concept/Skill", 30), ("Node text", 42), ("Action", 9),
    ("Standard", 14), ("Standard text", 48), ("Why", 40),
    ("Suggested by", 13), ("Suggested on", 13), ("Note", 30),
]
_EDITS_SHEET = "Ladder edits"
_HEADER_ROW = 5


def _font(**kw) -> Font:
    return Font(name=_FONT, size=kw.pop("size", 10), **kw)


def write_xlsx(edits: list[LadderEdit], exported_on: str,
               ladders_read: str | None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _EDITS_SHEET

    ws["A1"] = "Ladder edits to make"
    ws["A1"].font = _font(size=14, bold=True)
    read = ladders_read or "unknown"
    ws["A2"] = (f"Exported {exported_on}. Ladders last read into the tool {read}: "
                f"an edit made after that date is still listed until the next rebuild.")
    ws["A3"] = ("Fill in the Done column as you go. Nothing else needs editing. "
                "See the How to use tab.")
    for ref in ("A2", "A3"):
        ws[ref].font = _font(italic=True, color="595959")

    for c, (header, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=_HEADER_ROW, column=c, value=header)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _BOX
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[_HEADER_ROW].height = 30

    first = _HEADER_ROW + 1
    for r, e in enumerate(edits, start=first):
        values = ["", e.ladder_file, e.stem_name, e.seq, e.concept_skill, e.node_text,
                  e.action, e.as_written, e.standard_text, e.why, e.suggested_by,
                  e.suggested_on, ATTENTION_NOTE if e.needs_attention else None]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = _font()
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _BOX
        ws.cell(row=r, column=7).fill = _ADD_FILL if e.action == "Add" else _REMOVE_FILL
        ws.cell(row=r, column=7).font = _font(bold=True)
        ws.cell(row=r, column=8).font = _font(bold=True)
        if e.needs_attention:
            ws.cell(row=r, column=13).fill = _ATTENTION_FILL
    last = first + len(edits) - 1

    if not edits:
        ws.cell(row=first, column=1,
                value="No ladder edits outstanding.").font = _font(italic=True)
    else:
        done = DataValidation(type="list", formula1='"Yes"', allow_blank=True)
        ws.add_data_validation(done)
        done.add(f"A{first}:A{last}")
        last_col = ws.cell(row=_HEADER_ROW, column=len(_COLUMNS)).column_letter
        ws.auto_filter.ref = f"A{_HEADER_ROW}:{last_col}{last}"

    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{_HEADER_ROW}:{_HEADER_ROW}"

    _by_ladder_sheet(wb, edits, first, last)
    _how_to_sheet(wb)
    # openpyxl writes formulas without cached values; have Excel compute the
    # By ladder counts on open rather than show blanks.
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _by_ladder_sheet(wb: Workbook, edits: list[LadderEdit], first: int, last: int) -> None:
    sm = wb.create_sheet("By ladder")
    for c, (header, width) in enumerate(
            [("Ladder file", 60), ("Adds", 8), ("Removes", 10), ("Done", 8), ("Left", 8)],
            start=1):
        cell = sm.cell(row=1, column=c, value=header)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        sm.column_dimensions[cell.column_letter].width = width
    if not edits:
        return

    def rng(col):
        return f"'{_EDITS_SHEET}'!${col}${first}:${col}${last}"

    ladders = sorted({e.ladder_file or "" for e in edits})
    for r, ladder in enumerate(ladders, start=2):
        sm.cell(row=r, column=1, value=ladder)
        sm.cell(row=r, column=2, value=f'=COUNTIFS({rng("B")},A{r},{rng("G")},"Add")')
        sm.cell(row=r, column=3, value=f'=COUNTIFS({rng("B")},A{r},{rng("G")},"Remove")')
        sm.cell(row=r, column=4, value=f'=COUNTIFS({rng("B")},A{r},{rng("A")},"Yes")')
        sm.cell(row=r, column=5, value=f"=B{r}+C{r}-D{r}")
        for c in range(1, 6):
            sm.cell(row=r, column=c).font = _font()
    total = len(ladders) + 2
    sm.cell(row=total, column=1, value="Total").font = _font(bold=True)
    for c, col in zip(range(2, 6), "BCDE"):
        cell = sm.cell(row=total, column=c, value=f"=SUM({col}2:{col}{total - 1})")
        cell.font = _font(bold=True)


def _how_to_sheet(wb: Workbook) -> None:
    hw = wb.create_sheet("How to use")
    hw.column_dimensions["A"].width = 100
    lines = [
        ("How to use this list", True),
        ("Each row is one standard code to add to, or remove from, one node's "
         "'Notes related to Standards, Grade, or Leaf' cell in a ladder.", False),
        ("Rows are in ladder order: ladder file, then node number (the node's "
         "position in the document), so you can work top to bottom through each file.",
         False),
        ("To find a node, open the ladder file and search (Ctrl+F) for the start "
         "of the Node text.", False),
        ("For a Remove row, the Standard column shows the code as the ladder "
         "currently spells it.", False),
        ("Done: choose Yes once you've made the edit. The By ladder tab counts "
         "what's left.", False),
        ("You don't need to report back. The next rebuild reads the ladders and "
         "every edit that has landed drops off the next download.", False),
        ("A yellow Note means the node was reworded or removed after the edit "
         "was suggested, and has no node number.", False),
        ("Every row was entered by the lead in the review tool, and entering it "
         "counts as sign-off. If a row looks wrong, raise it with the lead and "
         "don't make the edit.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = hw.cell(row=i, column=1, value=text)
        cell.font = _font(size=12 if bold else 10, bold=bold)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
