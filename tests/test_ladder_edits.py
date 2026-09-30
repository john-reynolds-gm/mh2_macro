"""
mh2/ladder_edits.py and GET /api/export/ladder-edits.xlsx -- the ladder
keeper's export (docs/brief_ladder_edits_export.md).

Synthetic file-backed databases, same helpers as test_reconcile_review.py.
"""
import io
import sqlite3
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config  # noqa: E402
from mh2 import ladder_edits, reconcile_review, review_store  # noqa: E402
from mh2.normalize import alias_key  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_reconcile_review import _mh2_db, _node, _seq_db, _standard, _tag  # noqa: E402


def _place(con, node_id, source_file, seq, concept):
    con.execute("UPDATE nodes SET source_file=?, seq=?, concept_skill=? WHERE node_id=?",
                (source_file, seq, concept, node_id))


@pytest.fixture
def dbs(tmp_path):
    """Two ladders' worth of nodes:

      TIM-06  untagged                      <- open proposal for 2.MD.C.7 (Add)
      TIM-08  tagged 2.MD.C.7               <- incorrect_tag (Remove)
      TIM-09  tagged 2.MD.C7 (alias spelling of 2.MD.C.7)
                                            <- incorrect_tag (Remove, shown as written)
      TIM-12  tagged 4.MD.A.2               <- open proposal already landed (excluded)
      TIM-13  untagged                      <- incorrect_tag already removed (excluded)
                                               + partial_coverage review (excluded)
      FRA-12  untagged                      <- withdrawn proposal (excluded)
    """
    mh2 = _mh2_db(tmp_path / "mh2.db")
    for sid, g in (("2.MD.C.7", "2"), ("4.MD.A.2", "4"), ("TX.4.3D", "4")):
        _standard(mh2, sid, g)
    mh2.execute("UPDATE standards SET text='Tell time to five minutes.' "
                "WHERE standard_id='2.MD.C.7'")
    for nid, key, seq in (("TIM-06", "TIM:06", 6), ("TIM-08", "TIM:08", 8),
                          ("TIM-09", "TIM:09", 9), ("TIM-12", "TIM:12", 12),
                          ("TIM-13", "TIM:13", 13)):
        _node(mh2, nid, key, grades=("2",), stem_id="TIM")
        _place(mh2, nid, "Time.docx", seq, "Concept  with\nwhitespace")
    _node(mh2, "FRA-12", "FRA:12", grades=("4",), stem_id="FRA")
    _place(mh2, "FRA-12", "Fractions.docx", 12, "Compare")
    _tag(mh2, "TIM-08", "2.MD.C.7")
    _tag(mh2, "TIM-09", "2.MD.C7")
    mh2.execute("INSERT INTO standard_alias (alias_key, tier, standard_id) VALUES (?,?,?)",
                (alias_key("2.MD.C7", "punct"), "punct", "2.MD.C.7"))
    _tag(mh2, "TIM-12", "4.MD.A.2")
    mh2.execute("INSERT INTO ingest_log (run_id, ts, source_file, action)"
                " VALUES ('r', '2026-09-28 10:00:00', 'Time.docx', 'upsert_node')")
    mh2.commit()
    mh2.close()

    seq = _seq_db(tmp_path / "mh2_seq.db")
    review_store.create_tag_proposal(seq, "2.MD.C.7", "TIM:06", "TIM:06", "Debbie",
                                     ladder_file="Time.docx", rationale="untagged node")
    review_store.create_tag_proposal(seq, "4.MD.A.2", "TIM:12", "TIM:12", "Debbie",
                                     ladder_file="Time.docx")
    pid = review_store.create_tag_proposal(seq, "TX.4.3D", "FRA:12", "FRA:12", "Debbie",
                                           ladder_file="Fractions.docx")
    review_store.withdraw_tag_proposal(seq, pid)
    review_store.set_tag_review(seq, "2.MD.C.7", "TIM:08", "incorrect_tag", "Debbie",
                                note="durations, not clock time",
                                ladder_file_seen="Time.docx", node_text_seen="TIM:08")
    review_store.set_tag_review(seq, "2.MD.C.7", "TIM:09", "incorrect_tag", "Debbie",
                                ladder_file_seen="Time.docx", node_text_seen="TIM:09")
    review_store.set_tag_review(seq, "4.MD.A.2", "TIM:13", "incorrect_tag", "Debbie",
                                ladder_file_seen="Time.docx", node_text_seen="TIM:13")
    review_store.set_tag_review(seq, "2.MD.C.7", "TIM:13", "partial_coverage", "Debbie")
    seq.close()
    return tmp_path


def _build(tmp_path):
    mh2 = sqlite3.connect(f"file:{tmp_path / 'mh2.db'}?mode=ro", uri=True)
    seq = _seq_db(tmp_path / "mh2_seq.db")
    try:
        return ladder_edits.build_ladder_edits(mh2, seq)
    finally:
        mh2.close()
        seq.close()


def test_exports_only_outstanding_adds_and_removes_in_ladder_order(dbs):
    edits = _build(dbs)
    assert [(e.ladder_file, e.seq, e.action, e.standard_id) for e in edits] == [
        ("Time.docx", 6, "Add", "2.MD.C.7"),
        ("Time.docx", 8, "Remove", "2.MD.C.7"),
        ("Time.docx", 9, "Remove", "2.MD.C.7"),
    ]
    add = edits[0]
    assert add.stem_name == "TIM"
    assert add.concept_skill == "Concept with whitespace"
    assert add.standard_text == "Tell time to five minutes."
    assert add.why == "untagged node"
    assert add.suggested_by == "Debbie"
    assert len(add.suggested_on) == 10
    assert edits[1].why == "durations, not clock time"
    assert not any(e.needs_attention for e in edits)


def test_remove_row_shows_the_code_as_the_ladder_spells_it(dbs):
    remove = [e for e in _build(dbs) if e.seq == 9][0]
    assert remove.standard_id == "2.MD.C.7"
    assert remove.as_written == "2.MD.C7"


def test_duplicate_open_proposals_export_once_from_the_oldest(dbs):
    seq = _seq_db(dbs / "mh2_seq.db")
    review_store.create_tag_proposal(seq, "2.MD.C.7", "TIM:06", "TIM:06", "Other",
                                     ladder_file="Time.docx", rationale="again")
    seq.close()
    adds = [e for e in _build(dbs) if e.action == "Add"]
    assert len(adds) == 1
    assert adds[0].why == "untagged node"


def test_reworded_node_is_exported_flagged_from_its_snapshot(dbs):
    seq = _seq_db(dbs / "mh2_seq.db")
    review_store.create_tag_proposal(seq, "TX.4.3D", "FRA:gone", "Old wording",
                                     "Debbie", ladder_file="Fractions.docx")
    review_store.set_tag_review(seq, "4.MD.A.2", "TIM:gone", "incorrect_tag", "Debbie",
                                ladder_file_seen="Time.docx", node_text_seen="Old  time")
    seq.close()

    flagged = [e for e in _build(dbs) if e.needs_attention]
    assert {(e.action, e.ladder_file, e.stem_name, e.seq, e.node_text) for e in flagged} == {
        ("Add", "Fractions.docx", "FRA", None, "Old wording"),
        ("Remove", "Time.docx", "TIM", None, "Old time"),
    }


def test_needs_attention_proposal_stays_in_the_export(dbs):
    seq = _seq_db(dbs / "mh2_seq.db")
    pid = review_store.create_tag_proposal(seq, "TX.4.3D", "FRA:gone", "Old wording",
                                           "Debbie", ladder_file="Fractions.docx")
    review_store.update_proposal_state(seq, pid, "needs_attention")
    seq.close()
    assert any(e.standard_id == "TX.4.3D" and e.needs_attention for e in _build(dbs))


def test_removal_state_classifies_all_three_cases(dbs):
    from mh2 import node_lookup
    from mh2.coverage import build_tags_by_standard
    mh2 = sqlite3.connect(f"file:{dbs / 'mh2.db'}?mode=ro", uri=True)
    by_key = node_lookup.build_source_key_lookup(mh2)
    tags = build_tags_by_standard(mh2)
    mh2.close()

    def state(sid, key):
        return ladder_edits.removal_state({"standard_id": sid, "source_key": key},
                                          by_key, tags)[0]
    assert state("2.MD.C.7", "TIM:08") == ladder_edits.STILL_TAGGED
    assert state("2.MD.C.7", "TIM:09") == ladder_edits.STILL_TAGGED  # via alias
    assert state("4.MD.A.2", "TIM:13") == ladder_edits.REMOVED
    assert state("4.MD.A.2", "TIM:gone") == ladder_edits.NEEDS_ATTENTION


def test_reconcile_reports_removals_and_never_modifies_tag_review(dbs):
    seq = _seq_db(dbs / "mh2_seq.db")
    before = review_store.list_all_tag_reviews(seq)
    seq.close()

    reconcile_review.run(dbs / "mh2.db", dbs / "mh2_seq.db", dbs / "out")

    seq = _seq_db(dbs / "mh2_seq.db")
    assert review_store.list_all_tag_reviews(seq) == before
    seq.close()
    report = (dbs / "out" / "review_reconcile.txt").read_text()
    removed = report.split("incorrect_tag: removed from the ladder (done)", 1)[1]
    assert removed.splitlines()[1] == "1 rows"
    assert "TIM:13" in removed.split("===", 2)[1]
    still = report.split("incorrect_tag: still tagged (outstanding work)", 1)[1]
    assert still.splitlines()[1] == "2 rows"


# ------------------------------------------------------------------ workbook

def test_workbook_rows_match_edits(dbs):
    edits = _build(dbs)
    wb = load_workbook(io.BytesIO(
        ladder_edits.write_xlsx(edits, exported_on="2026-09-30", ladders_read="2026-09-28")))
    assert wb.sheetnames == ["Ladder edits", "By ladder", "How to use"]
    ws = wb["Ladder edits"]
    assert "2026-09-28" in ws["A2"].value
    header = [c.value for c in ws[5]]
    assert header[:8] == ["Done", "Ladder file", "Stem", "Node #", "Concept/Skill",
                          "Node text", "Action", "Standard"]
    body = [[c.value for c in row] for row in ws.iter_rows(min_row=6)]
    assert [(r[1], r[3], r[6], r[7]) for r in body] == [
        ("Time.docx", 6, "Add", "2.MD.C.7"),
        ("Time.docx", 8, "Remove", "2.MD.C.7"),
        ("Time.docx", 9, "Remove", "2.MD.C7"),
    ]
    assert ws.auto_filter.ref == "A5:M8"
    by_ladder = wb["By ladder"]
    assert by_ladder["A2"].value == "Time.docx"
    assert by_ladder["B2"].value.startswith("=COUNTIFS('Ladder edits'!$B$6:$B$8,A2,")
    assert by_ladder["A3"].value == "Total"


def test_workbook_with_nothing_outstanding(tmp_path):
    wb = load_workbook(io.BytesIO(
        ladder_edits.write_xlsx([], exported_on="2026-09-30", ladders_read=None)))
    ws = wb["Ladder edits"]
    assert ws["A6"].value == "No ladder edits outstanding."
    assert "unknown" in ws["A2"].value
    assert wb["By ladder"].max_row == 1


# ----------------------------------------------------------------------- API

def test_export_endpoint_downloads_the_workbook(dbs, monkeypatch):
    monkeypatch.setattr(config, "DB", dbs / "mh2.db")
    monkeypatch.setattr(config, "SEQ_DB", dbs / "mh2_seq.db")
    import review_api
    resp = TestClient(review_api.app).get("/api/export/ladder-edits.xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="ladder_edits_' in resp.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(resp.content))["Ladder edits"]
    assert ws.max_row == 8
    assert "2026-09-28" in ws["A2"].value
