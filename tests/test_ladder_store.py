"""
mh2/ladder_store.py -- writer-uploaded ladders (docs/brief_ladder_uploads.md).
Numbers below are the brief's §9 acceptance criteria.

  #1  validation refuses non-docx, '~$' lock files, files without the PK
      magic, oversized files, zero-node documents, and unknown stem_ids
  #2  an upload to a DRAFTED slot is ingested under the slot's name, and the
      effective ladder dir holds the upload, not the repo copy
  #3  an upload to an UNDRAFTED slot registers it: the effective stems.csv has
      ladder_drafted=1 on every row of that stem_id, and
      load_stems.resolve_ladder_files accepts it
  #4  one live version per stem; revert flips it
  #5  diff counts for an edited ladder -- one reworded node is 1 new +
      1 removed -- and the orphan count reflects a seeded review decision
  #6  no uploads -> the effective inputs are the repo inputs (regression);
      --no-uploads ignores uploads that exist
  #8  uploads and their metadata survive deleting mh2.db and rerunning the
      entrypoint's rebuild step

(#7 and #9 are HTTP-level and live in test_ladder_upload_api.py.)

Ladders here are tiny .docx files built with python-docx in the shape
ingest_ladders.read_docx reads: a 'Concept/Skill:' paragraph followed by a
table whose 'Knowledge Graph Node' row names one node per column.
"""
import csv
import io
import os
import sqlite3
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import config  # noqa: E402
import rebuild  # noqa: E402
from mh2 import ladder_store, review_store  # noqa: E402
from mh2.ingest_ladders import source_key  # noqa: E402
from mh2.load_stems import read_stems_csv, resolve_ladder_files  # noqa: E402
from test_reconcile_review import _mh2_db, _seq_db  # noqa: E402

DRAFTED = "NS-COUNT"
# Declared with underscores, on disk with spaces -- the normalization case six
# of the real drafted stems are in.
DRAFTED_DECLARED = "MH2_PK5_Counting_Things.docx"
DRAFTED_ON_DISK = "MH2_PK5_Counting Things.docx"
OTHER = "NS-FRAC"
OTHER_FILE = "MH2_PK5_Fractions.docx"
UNDRAFTED = "OE-FOO"          # two workbook stems under one stem_id
REAL_DATA = config.DATA       # captured before any test repoints config
NODES = ["Count to 10", "Count to 20", "Count to 100"]

FIELDS = ("stem_id", "band", "stem_group", "masterlist_name", "workbook_sheet",
          "workbook_stem", "ladder_file", "ladder_drafted")
STEMS_ROWS = [
    (DRAFTED, "PK5", "Number Systems", "Counting", "NSS", "Counting", DRAFTED_DECLARED, "1"),
    (OTHER, "PK5", "Number Systems", "Fractions", "NSS", "Fractions", OTHER_FILE, "1"),
    (UNDRAFTED, "PK5", "Operations", "Foo and Bar", "OE", "Foo", "", "0"),
    (UNDRAFTED, "PK5", "Operations", "Foo and Bar", "OE", "Bar", "", "0"),
]


def make_ladder(path: Path, nodes: list[str], heading: str = "Count things") -> Path:
    from docx import Document
    doc = Document()
    doc.add_paragraph(f"Concept/Skill: {heading}")
    if nodes:
        t = doc.add_table(rows=2, cols=len(nodes) + 1)
        t.cell(0, 0).text = "Knowledge Graph Node"
        t.cell(1, 0).text = "Grade (or Leaf)"
        for i, n in enumerate(nodes):
            t.cell(0, i + 1).text = n
            t.cell(1, i + 1).text = "K"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def docx_bytes(nodes: list[str], tmp: Path, name="upload.docx") -> bytes:
    return make_ladder(tmp / "_make" / name, nodes).read_bytes()


def write_stems_csv(path: Path, rows=STEMS_ROWS) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(FIELDS)
        w.writerows(rows)
    return path


def point_config_at(monkeypatch, data: Path) -> None:
    """Repoint every config path ladder_store and review_api read at a data
    tree laid out the way config.py lays out the real one."""
    build = data / "build"
    monkeypatch.setattr(config, "DATA", data)
    monkeypatch.setattr(config, "SOURCE", data / "source")
    monkeypatch.setattr(config, "BUILD", build)
    monkeypatch.setattr(config, "REPORTS", data / "reports")
    monkeypatch.setattr(config, "LADDERS", data / "source" / "ladders")
    monkeypatch.setattr(config, "STEMS_CSV", data / "source" / "workbooks" / "stems.csv")
    monkeypatch.setattr(config, "DB", build / "mh2.db")
    monkeypatch.setattr(config, "SEQ_DB", build / "mh2_seq.db")
    monkeypatch.setattr(config, "LADDER_UPLOADS", build / "ladder_uploads")
    monkeypatch.setattr(config, "EFFECTIVE_INPUTS", build / "_effective")


def seed_tree(data: Path) -> Path:
    """Repo ladders (two drafted, one Word lock file), stems.csv, an empty
    mh2_seq.db, and an mh2.db serving NODES for the drafted stem."""
    lad = data / "source" / "ladders"
    make_ladder(lad / DRAFTED_ON_DISK, NODES)
    make_ladder(lad / OTHER_FILE, ["Name a unit fraction"], "Fractions")
    (lad / ("~$" + DRAFTED_ON_DISK[2:])).write_bytes(b"lock")
    write_stems_csv(data / "source" / "workbooks" / "stems.csv")
    for d in ("build", "reports"):
        (data / d).mkdir(parents=True, exist_ok=True)
    _seq_db(data / "build" / "mh2_seq.db").close()
    mh2 = _mh2_db(data / "build" / "mh2.db")
    mh2.execute("INSERT INTO stems (stem_id, name) VALUES (?,?)", (DRAFTED, "Counting"))
    for i, text in enumerate(NODES, 1):
        mh2.execute("INSERT INTO nodes (node_id, stem_id, seq, source_key, node_text,"
                    " source_file) VALUES (?,?,?,?,?,?)",
                    (f"{DRAFTED}-{i:04d}", DRAFTED, i, source_key(DRAFTED, text), text,
                     DRAFTED_ON_DISK))
    mh2.commit()
    mh2.close()
    return data


@pytest.fixture
def tree(tmp_path, monkeypatch):
    data = seed_tree(tmp_path / "data")
    point_config_at(monkeypatch, data)
    con = ladder_store.connect()
    yield data, con
    con.close()


def validate(con, stem_id, data: bytes, name="My Ladder FINAL (2).docx", by="jane"):
    return ladder_store.validate_upload(con, stem_id, name, io.BytesIO(data), by)


def upload(con, tmp, stem_id, nodes, by="jane", note=None, confirm=False):
    p = validate(con, stem_id, docx_bytes(nodes, tmp), by=by)
    return ladder_store.commit_upload(con, stem_id, p["token"], by, note=note,
                                      confirm_orphans=confirm)


# ----------------------------------------------------------------------- #1

@pytest.mark.parametrize("name,data,code", [
    ("ladder.doc", b"PK\x03\x04whatever", "not_docx"),
    ("ladder.pdf", b"%PDF-1.4", "not_docx"),
    ("~$_PK5_Counting.docx", b"PK\x03\x04", "lock_file"),
    ("ladder.docx", b"\xd0\xcf\x11\xe0 old binary .doc renamed", "not_docx"),
    ("ladder.docx", b"", "empty_file"),
    ("ladder.docx", b"PK\x03\x04 but not a real zip", "unreadable"),
])
def test_validation_refuses_bad_files(tree, name, data, code):
    _, con = tree
    with pytest.raises(ladder_store.LadderUploadError) as ei:
        validate(con, DRAFTED, data, name=name)
    assert ei.value.status == 422
    assert ei.value.code == code


def test_validation_refuses_oversized_files(tree, monkeypatch):
    _, con = tree
    monkeypatch.setattr(ladder_store, "MAX_BYTES", 1024)
    with pytest.raises(ladder_store.LadderUploadError) as ei:
        validate(con, DRAFTED, b"PK" + b"\0" * 4096)
    assert ei.value.code == "too_large"


def test_validation_refuses_a_document_with_no_nodes(tree, tmp_path):
    _, con = tree
    with pytest.raises(ladder_store.LadderUploadError) as ei:
        validate(con, DRAFTED, docx_bytes([], tmp_path))
    assert ei.value.code == "no_nodes"


def test_validation_refuses_an_unknown_stem_without_recording_it(tree, tmp_path):
    _, con = tree
    with pytest.raises(ladder_store.LadderUploadError) as ei:
        validate(con, "NOT-A-STEM", docx_bytes(NODES, tmp_path))
    assert ei.value.status == 404
    assert ladder_store.list_rejections(con) == {}


def test_a_refusal_is_recorded_as_the_stems_rejected_status_and_leaves_no_file(tree):
    data, con = tree
    with pytest.raises(ladder_store.LadderUploadError):
        validate(con, DRAFTED, b"nope", name="ladder.docx")
    rej = ladder_store.get_rejection(con, DRAFTED)
    assert rej["attempted_by"] == "jane" and rej["original_name"] == "ladder.docx"
    row = {r["stem_id"]: r for r in ladder_store.list_slots(con)}[DRAFTED]
    assert row["status"] == "rejected" and row["source"] == "repo"
    assert list(ladder_store.pending_dir().iterdir()) == []


def test_a_successful_commit_clears_the_rejection(tree, tmp_path):
    _, con = tree
    with pytest.raises(ladder_store.LadderUploadError):
        validate(con, DRAFTED, b"nope", name="ladder.docx")
    upload(con, tmp_path, DRAFTED, NODES)
    assert ladder_store.get_rejection(con, DRAFTED) is None


# ----------------------------------------------------------------------- #2

def test_drafted_upload_lands_in_the_effective_dir_instead_of_the_repo_copy(tree, tmp_path):
    data, con = tree
    row = upload(con, tmp_path, DRAFTED, NODES + ["Count to 1000"])
    stored = ladder_store.version_path(row)
    assert stored.exists() and stored.parent.name == DRAFTED
    assert row["original_name"] == "My Ladder FINAL (2).docx"

    ladders, stems = ladder_store.assemble_effective_inputs(config.EFFECTIVE_INPUTS)
    names = sorted(p.name for p in ladders.iterdir())
    # The repo's on-disk spelling, not the writer's filename and not the
    # declared underscore spelling; exactly one file for the stem; no lock file.
    assert names == sorted([DRAFTED_ON_DISK, OTHER_FILE])
    assert (ladders / DRAFTED_ON_DISK).resolve() == stored.resolve()
    assert (ladders / OTHER_FILE).resolve() == (data / "source/ladders" / OTHER_FILE).resolve()
    # A drafted slot needs no registration; stems.csv passes through untouched.
    assert ladder_store.list_overrides(con) == {}
    assert stems.read_bytes() == config.STEMS_CSV.read_bytes()
    assert resolve_ladder_files(read_stems_csv(stems), ladders)[DRAFTED].resolve() \
        == stored.resolve()


def test_repo_ladders_are_never_modified(tree, tmp_path):
    data, con = tree
    before = {p.name: p.read_bytes() for p in (data / "source/ladders").iterdir()}
    upload(con, tmp_path, DRAFTED, ["Something else entirely"])
    ladder_store.assemble_effective_inputs(config.EFFECTIVE_INPUTS)
    after = {p.name: p.read_bytes() for p in (data / "source/ladders").iterdir()}
    assert before == after


# ----------------------------------------------------------------------- #3

def test_undrafted_upload_registers_every_row_of_the_stem(tree, tmp_path):
    _, con = tree
    row = upload(con, tmp_path, UNDRAFTED, ["Foo one", "Bar one"])
    override = ladder_store.get_override(con, UNDRAFTED)
    assert override["ladder_file"] == f"MH2_PK5_{UNDRAFTED}.docx"
    assert override["set_by"] == "jane"

    ladders, stems = ladder_store.assemble_effective_inputs(config.EFFECTIVE_INPUTS)
    rows = [r for r in read_stems_csv(stems) if r["stem_id"] == UNDRAFTED]
    assert len(rows) == 2
    assert all(r["ladder_drafted"] == 1 and r["ladder_file"] == override["ladder_file"]
               for r in rows)
    resolved = resolve_ladder_files(read_stems_csv(stems), ladders)
    assert resolved[UNDRAFTED].name == override["ladder_file"]
    assert resolved[UNDRAFTED].resolve() == ladder_store.version_path(row).resolve()
    # Every other row is written back verbatim.
    orig = list(csv.reader(open(config.STEMS_CSV, encoding="utf-8")))
    eff = list(csv.reader(open(stems, encoding="utf-8")))
    assert [r for r in eff if r[0] != UNDRAFTED] == [r for r in orig if r[0] != UNDRAFTED]


def test_ingest_accepts_the_effective_inputs_for_an_undrafted_upload(tree, tmp_path):
    """The real ingest step, against the assembled inputs: the override is
    what lets ingest_ladders accept a file stems.csv never declared."""
    data, con = tree
    upload(con, tmp_path, UNDRAFTED, ["Foo one", "Bar one"])
    ladders, stems = ladder_store.assemble_effective_inputs(config.EFFECTIVE_INPUTS)
    db = tmp_path / "ingest.db"
    _mh2_db(db).close()
    res = subprocess.run(
        [sys.executable, "-m", "mh2.ingest_ladders", "--db", str(db),
         "--glob", str(ladders / "*.docx"), "--csv", str(stems)],
        cwd=REPO, capture_output=True, text=True)
    assert res.returncode == 0, res.stdout + res.stderr
    con2 = sqlite3.connect(db)
    got = dict(con2.execute("SELECT stem_id, COUNT(*) FROM nodes GROUP BY stem_id"))
    con2.close()
    assert got == {DRAFTED: 3, OTHER: 1, UNDRAFTED: 2}


def test_a_stems_csv_row_added_later_wins_the_name_over_the_override(tree, tmp_path, monkeypatch):
    """If John later commits a ladder_file for a stem that was registered by
    upload, the repo declaration names the file; the upload still wins the
    content (§3)."""
    data, con = tree
    row = upload(con, tmp_path, UNDRAFTED, ["Foo one"])
    rows = [r if r[0] != UNDRAFTED else r[:6] + ("MH2_PK5_Foo.docx", "1") for r in STEMS_ROWS]
    write_stems_csv(config.STEMS_CSV, rows)
    make_ladder(data / "source/ladders/MH2_PK5_Foo.docx", ["Repo foo"])
    ladders, stems = ladder_store.assemble_effective_inputs(config.EFFECTIVE_INPUTS)
    assert (ladders / "MH2_PK5_Foo.docx").resolve() == ladder_store.version_path(row).resolve()
    assert not (ladders / f"MH2_PK5_{UNDRAFTED}.docx").exists()
    assert stems.read_bytes() == config.STEMS_CSV.read_bytes()


# ----------------------------------------------------------------------- #4

def test_only_one_live_version_and_revert_flips_it(tree, tmp_path):
    _, con = tree
    first = upload(con, tmp_path, DRAFTED, NODES, note="v1")
    second = upload(con, tmp_path, DRAFTED, NODES + ["Count to 1000"], confirm=True)
    live = [v["upload_id"] for v in ladder_store.list_versions(con, DRAFTED) if v["is_live"]]
    assert live == [second["upload_id"]]

    ladder_store.revert(con, DRAFTED, first["upload_id"])
    live = [v["upload_id"] for v in ladder_store.list_versions(con, DRAFTED) if v["is_live"]]
    assert live == [first["upload_id"]]
    ladders, _ = ladder_store.assemble_effective_inputs(config.EFFECTIVE_INPUTS)
    assert (ladders / DRAFTED_ON_DISK).resolve() == ladder_store.version_path(first).resolve()
    # Both files are still on disk: nothing is overwritten or deleted.
    assert ladder_store.version_path(second).exists()


def test_the_database_itself_refuses_two_live_versions(tree, tmp_path):
    _, con = tree
    upload(con, tmp_path, DRAFTED, NODES)
    with pytest.raises(sqlite3.IntegrityError):
        con.execute(
            "INSERT INTO ladder_upload (stem_id, stored_path, original_name, sha256,"
            " size_bytes, uploaded_by, uploaded_at, node_count, is_live)"
            " VALUES (?,?,?,?,?,?,?,?,1)", (DRAFTED, "x", "x", "x", 1, "x", "x", 1))


def test_revert_refuses_another_stems_version(tree, tmp_path):
    _, con = tree
    other = upload(con, tmp_path, OTHER, ["Name a unit fraction"])
    with pytest.raises(ladder_store.LadderUploadError) as ei:
        ladder_store.revert(con, DRAFTED, other["upload_id"])
    assert ei.value.status == 404


def test_commit_refuses_a_token_validated_for_another_stem(tree, tmp_path):
    _, con = tree
    p = validate(con, OTHER, docx_bytes(["Name a unit fraction"], tmp_path))
    with pytest.raises(ladder_store.LadderUploadError) as ei:
        ladder_store.commit_upload(con, DRAFTED, p["token"], "jane")
    assert ei.value.code == "wrong_stem"


@pytest.mark.parametrize("token", ["../../etc/passwd", "abc", ""])
def test_commit_refuses_a_malformed_token(tree, token):
    _, con = tree
    with pytest.raises(ladder_store.LadderUploadError) as ei:
        ladder_store.commit_upload(con, DRAFTED, token, "jane")
    assert ei.value.code == "bad_token"


def test_a_token_is_single_use(tree, tmp_path):
    _, con = tree
    p = validate(con, DRAFTED, docx_bytes(NODES, tmp_path))
    ladder_store.commit_upload(con, DRAFTED, p["token"], "jane")
    with pytest.raises(ladder_store.LadderUploadError) as ei:
        ladder_store.commit_upload(con, DRAFTED, p["token"], "jane")
    assert ei.value.status == 410


# ----------------------------------------------------------------------- #5

def test_diff_counts_one_reworded_node_as_one_new_and_one_removed(tree, tmp_path):
    _, con = tree
    edited = [NODES[0], "Count to 20 by ones", NODES[2]]
    p = validate(con, DRAFTED, docx_bytes(edited, tmp_path))
    assert p["node_count"] == 3
    assert p["diff"]["matched"] == 2
    assert p["diff"]["new"] == 1
    assert p["diff"]["removed"] == 1
    assert p["diff"]["removed_texts"] == [NODES[1]]
    assert p["orphaned_review_count"] == 0


def test_orphan_count_reflects_review_decisions_on_removed_nodes(tree, tmp_path):
    """Seeds one tag_review and one open tag_proposal on the reworded node,
    plus a withdrawn proposal (not counted) and a review on a node that
    survives (not counted): the literal expected count is 2."""
    _, con = tree
    removed_key = source_key(DRAFTED, NODES[1])
    kept_key = source_key(DRAFTED, NODES[0])
    review_store.set_tag_review(con, "K.CC.A.1", removed_key, "confirmed", "amanda")
    review_store.set_tag_review(con, "K.CC.A.2", kept_key, "confirmed", "amanda")
    review_store.create_tag_proposal(con, "K.CC.A.3", removed_key, NODES[1], "amanda")
    pid = review_store.create_tag_proposal(con, "K.CC.A.4", removed_key, NODES[1], "amanda")
    review_store.withdraw_tag_proposal(con, pid, None)

    p = validate(con, DRAFTED, docx_bytes([NODES[0], "Count to 20 by ones", NODES[2]], tmp_path))
    assert p["orphans"] == {"tag_reviews": 1, "open_proposals": 1,
                            "candidate_rulings": 0, "placements": 0}
    assert p["orphaned_review_count"] == 2

    with pytest.raises(ladder_store.LadderUploadError) as ei:
        ladder_store.commit_upload(con, DRAFTED, p["token"], "jane")
    assert ei.value.status == 409 and ei.value.code == "confirm_orphans"
    row = ladder_store.commit_upload(con, DRAFTED, p["token"], "jane", confirm_orphans=True)
    assert row["is_live"] == 1


def test_diff_against_a_stem_with_no_live_nodes_is_all_new(tree, tmp_path):
    _, con = tree
    p = validate(con, UNDRAFTED, docx_bytes(["Foo one", "Bar one"], tmp_path))
    assert p["diff"] == {"matched": 0, "new": 2, "removed": 0, "live_node_count": 0,
                         "removed_texts": []}


def test_same_as_live_is_flagged(tree, tmp_path):
    _, con = tree
    data = docx_bytes(NODES, tmp_path)
    p = validate(con, DRAFTED, data)
    ladder_store.commit_upload(con, DRAFTED, p["token"], "jane")
    assert validate(con, DRAFTED, data)["same_as_live"] is True


# ----------------------------------------------------------------------- #6

def test_with_no_uploads_the_effective_inputs_are_the_repo_inputs(tree):
    """The regression half of #6 that runs in the suite. The database half
    was checked once by hand against the real data/source: HEAD's rebuild.py
    and this one, each into its own MH2_DATA_DIR, produced the same rows in
    all 30 tables except run timestamps, ingest_log.run_id (a hash of the
    glob path) and stem_map.ladder_path (now under _effective; its one
    reader, mh2/reconcile.py, uses only the file name)."""
    data, _ = tree
    ladders, stems = ladder_store.assemble_effective_inputs(config.EFFECTIVE_INPUTS)
    repo = {p.name: p for p in (data / "source/ladders").iterdir()
            if not p.name.startswith("~$")}
    eff = {p.name: p for p in ladders.iterdir()}
    assert sorted(eff) == sorted(repo)
    assert all(eff[n].read_bytes() == repo[n].read_bytes() for n in repo)
    assert stems.read_bytes() == config.STEMS_CSV.read_bytes()


def test_effective_dir_is_rebuilt_from_scratch(tree, tmp_path):
    _, con = tree
    stale = config.EFFECTIVE_INPUTS / "ladders" / "MH2_PK5_Stale.docx"
    stale.parent.mkdir(parents=True)
    stale.write_bytes(b"old")
    ladders, _ = ladder_store.assemble_effective_inputs(config.EFFECTIVE_INPUTS)
    assert not stale.exists()


def test_a_live_upload_for_a_stem_dropped_from_stems_csv_is_left_out(tree, tmp_path, capsys):
    _, con = tree
    upload(con, tmp_path, UNDRAFTED, ["Foo one"])
    write_stems_csv(config.STEMS_CSV, [r for r in STEMS_ROWS if r[0] != UNDRAFTED])
    ladders, stems = ladder_store.assemble_effective_inputs(config.EFFECTIVE_INPUTS)
    assert sorted(p.name for p in ladders.iterdir()) == sorted([DRAFTED_ON_DISK, OTHER_FILE])
    assert "no longer in stems.csv" in capsys.readouterr().out


def _rebuild_tree(data: Path) -> Path:
    """Every input rebuild.py's pre-flight requires, present but empty except
    stems.csv and the ladders -- the run clears pre-flight, assembles the
    effective inputs, and then fails in the first loader in about a second
    (the same shape test_rebuild_hardening.py relies on)."""
    for p in rebuild.REQUIRED_INPUTS:
        dest = data / p.relative_to(REAL_DATA)
        dest.parent.mkdir(parents=True, exist_ok=True)
        if not dest.exists():
            dest.write_bytes(b"")
    return data


def run_rebuild(data: Path, *flags: str) -> subprocess.CompletedProcess:
    env = dict(os.environ, MH2_DATA_DIR=str(data), MH2_REBUILD_ACTOR="tester")
    env.pop("MH2_LADDERS_DIR", None)
    env.pop("MH2_STEMS_CSV", None)
    return subprocess.run(
        [sys.executable, str(REPO / "scripts" / "rebuild.py"), "--skip-review-reconcile", *flags],
        capture_output=True, text=True, cwd=str(REPO), env=env, timeout=300)


@pytest.fixture
def rebuild_tree(tmp_path, monkeypatch):
    data = tmp_path / "data"
    point_config_at(monkeypatch, data)
    seed_tree(data)
    _rebuild_tree(data)
    con = ladder_store.connect()
    yield data, con
    con.close()


def test_rebuild_reads_uploads_and_no_uploads_ignores_them(rebuild_tree, tmp_path):
    data, con = rebuild_tree
    row = upload(con, tmp_path, DRAFTED, ["Uploaded node"])

    res = run_rebuild(data, "--no-uploads")
    assert "repo copies only" in res.stdout
    assert not (data / "build/_effective").exists()

    res = run_rebuild(data)
    assert "assembling ladder inputs" in res.stdout
    eff = data / "build/_effective/ladders" / DRAFTED_ON_DISK
    assert eff.resolve() == ladder_store.version_path(row).resolve()


# ----------------------------------------------------------------------- #8

def test_uploads_survive_deleting_mh2_db_and_rerunning_the_rebuild(rebuild_tree, tmp_path):
    """entrypoint.sh's first step is `rebuild.py --skip-review-reconcile`.
    Deleting mh2.db and running it again must leave every upload file and
    every ladder_upload / stem_registry_override row in place, and the next
    build must still read the uploads."""
    data, con = rebuild_tree
    drafted = upload(con, tmp_path, DRAFTED, ["Uploaded node"])
    undrafted = upload(con, tmp_path, UNDRAFTED, ["Foo one"])
    before_rows = [dict(r) for r in con.execute("SELECT * FROM ladder_upload ORDER BY upload_id")]
    before_override = ladder_store.list_overrides(con)
    before_files = {p: p.read_bytes() for p in config.LADDER_UPLOADS.rglob("*.docx")}
    con.close()

    (data / "build/mh2.db").unlink()
    run_rebuild(data)
    run_rebuild(data)

    con2 = ladder_store.connect()
    try:
        assert [dict(r) for r in con2.execute(
            "SELECT * FROM ladder_upload ORDER BY upload_id")] == before_rows
        assert ladder_store.list_overrides(con2) == before_override
    finally:
        con2.close()
    assert {p: p.read_bytes() for p in config.LADDER_UPLOADS.rglob("*.docx")} == before_files
    eff = data / "build/_effective"
    assert (eff / "ladders" / DRAFTED_ON_DISK).resolve() == \
        ladder_store.version_path(drafted).resolve()
    assert (eff / "ladders" / f"MH2_PK5_{UNDRAFTED}.docx").resolve() == \
        ladder_store.version_path(undrafted).resolve()


# ------------------------------------------------------------- housekeeping

def test_purge_pending_drops_only_files_older_than_a_day(tree, tmp_path):
    _, con = tree
    fresh = validate(con, DRAFTED, docx_bytes(NODES, tmp_path))["token"]
    stale = validate(con, DRAFTED, docx_bytes(NODES, tmp_path))["token"]
    old = 2 * 24 * 3600
    for ext in (".docx", ".json"):
        p = ladder_store.pending_dir() / f"{stale}{ext}"
        st = p.stat()
        os.utime(p, (st.st_atime - old, st.st_mtime - old))
    assert ladder_store.purge_pending() == 2
    left = sorted(p.name for p in ladder_store.pending_dir().iterdir())
    assert left == sorted([f"{fresh}.docx", f"{fresh}.json"])


def test_list_slots_reports_source_and_status_per_stem(tree, tmp_path):
    data, con = tree
    upload(con, tmp_path, DRAFTED, NODES)
    (data / "source/ladders" / OTHER_FILE).unlink()
    rows = {r["stem_id"]: r for r in ladder_store.list_slots(con)}
    assert list(rows) == [DRAFTED, OTHER, UNDRAFTED]       # one per stem_id
    assert (rows[DRAFTED]["source"], rows[DRAFTED]["status"]) == ("upload", "uploaded")
    assert rows[DRAFTED]["live"]["uploaded_by"] == "jane"
    assert rows[DRAFTED]["version_count"] == 1
    assert (rows[OTHER]["source"], rows[OTHER]["status"]) == ("none", "empty")
    assert (rows[UNDRAFTED]["source"], rows[UNDRAFTED]["status"]) == ("none", "empty")
    assert rows[UNDRAFTED]["workbook_stems"] == ["Foo", "Bar"]
    assert rows[UNDRAFTED]["canonical_name"] == f"MH2_PK5_{UNDRAFTED}.docx"
